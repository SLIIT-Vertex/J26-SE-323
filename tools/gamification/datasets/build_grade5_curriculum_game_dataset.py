from __future__ import annotations

import argparse
import hashlib
import json
import os
import random
import re
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path


PROJECT_ROOT = Path(__file__).resolve().parents[3]
DEFAULT_OUTPUT_DIR = (
    PROJECT_ROOT / "data" / "gamification" / "grade5-curriculum-game-dataset" / "v2"
)
DEFAULT_CURRICULUM_OUTPUT_DIR = (
    PROJECT_ROOT / "data" / "curriculum" / "grade5-curriculum-trace" / "v1"
)
CURRICULUM_TRACE_PACKAGE = "grade5-curriculum-trace"
CURRICULUM_TRACE_VERSION = "1.0.0"
CURRICULUM_TRACE_FILES = (
    "curriculum_sources.jsonl",
    "curriculum_nodes.jsonl",
    "curriculum_terms.jsonl",
)
SOURCE_DIR_ENV = "MINDBRIDGE_GRADE5_SYLLABUS_DIR"
PACKAGE_VERSION = "2.0.0"
CURRICULUM_VERSION = "grade5-curriculum-v2"
SKILL_TAXONOMY_VERSION = "grade5-skill-candidates-v2"
MISCONCEPTION_CATALOGUE_VERSION = "grade5-misconceptions-v2"
SUPPORT_CATALOGUE_VERSION = "grade5-support-v2"
EQUIVALENCE_VERSION = "grade5-verification-candidates-v2"
MASTER_WORKBOOK_NAME = "grade5_game_dataset_master.xlsx"
SUBJECTS = ["ENV", "SIN", "ENG", "MAT"]
SUBJECT_NAMES = {
    "ENV": "Environmental Studies",
    "SIN": "Sinhala",
    "ENG": "English",
    "MAT": "Mathematics",
}
LANGUAGES = {"ENV": "si", "SIN": "si", "ENG": "en", "MAT": "bilingual"}
RNG = random.Random(5323)


def clean(value):
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def json_value(value):
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, (bool, int, float)):
        return value
    return clean(value)


def source_record(row):
    return {
        key: json_value(value)
        for key, value in row.items()
        if not key.startswith("_")
    }


def read_sheet(workbook, sheet_name):
    ws = workbook[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    headers = [clean(v) for v in rows[0]]
    records = []
    for row_number, row in enumerate(rows[1:], start=2):
        if not any(v is not None for v in row):
            continue
        record = dict(zip(headers, row))
        record["_source_row_number"] = row_number
        records.append(record)
    return records


def first(row, *keys):
    for key in keys:
        value = clean(row.get(key))
        if value:
            return value
    return ""


def derive_parent_from_id(node_id, node_type):
    if node_type == "learning_outcome":
        return re.sub(r"-O\d+$", "", node_id)
    if node_type == "content_point":
        return re.sub(r"-K\d+$", "", node_id)
    return ""


def jsonl_write(path, rows):
    with path.open("w", encoding="utf-8") as handle:
        for row in rows:
            handle.write(json.dumps(row, ensure_ascii=False) + "\n")


def fallback_distractors(correct):
    text = clean(correct)
    time_match = re.fullmatch(r"(\d{2}):(\d{2})", text)
    if time_match:
        total_minutes = int(time_match.group(1)) * 60 + int(time_match.group(2))
        return [
            f"{((total_minutes + delta) // 60) % 24:02d}:{(total_minutes + delta) % 60:02d}"
            for delta in (15, -15, 30, -30, 60, -60)
        ]
    numeric_match = re.fullmatch(r"(-?\d+)(.*)", text)
    if numeric_match:
        number = int(numeric_match.group(1))
        suffix = numeric_match.group(2)
        deltas = (1, -1, 2, -2, 10, -10, 100, -100)
        return [f"{max(0, number + delta)}{suffix}" for delta in deltas]
    decimal_match = re.fullmatch(r"(-?\d+\.\d+)", text)
    if decimal_match:
        number = float(decimal_match.group(1))
        return [f"{max(0, number + delta):.1f}" for delta in (0.1, -0.1, 0.2, -0.2)]
    return []


def make_options(correct, pool, count=4):
    candidates = list(
        dict.fromkeys(
            clean(v) for v in pool if clean(v) and clean(v) != clean(correct)
        )
    )
    RNG.shuffle(candidates)
    for fallback in fallback_distractors(correct):
        if len(candidates) >= count - 1:
            break
        if fallback != clean(correct) and fallback not in candidates:
            candidates.append(fallback)
    options = [clean(correct)] + candidates[: max(0, count - 1)]
    if len(options) < count:
        raise ValueError(f"not enough unique distractors for {correct!r}")
    RNG.shuffle(options)
    return options


def source_confidence(row):
    return first(row, "extraction_confidence") or "not_stated"


def collect_curriculum(source_dir: Path):
    from openpyxl import load_workbook

    nodes = []
    raw = {}
    for subject in SUBJECTS:
        path = source_dir / f"grade5_{subject}_syllabus_trace_v1.xlsx"
        wb = load_workbook(path, read_only=True, data_only=True)
        subject_raw = {}
        for sheet in wb.sheetnames:
            subject_raw[sheet] = read_sheet(wb, sheet)
        raw[subject] = subject_raw

        node_specs = [
            ("02_competencies", "competency", "competency_id", "", ("competency", "competency_si"), ("competency", "competency_en_gloss")),
            ("03_competency_levels", "competency_level", "competency_level_id", "competency_id", ("competency_level", "competency_level_si"), ("competency_level", "competency_level_en_gloss")),
            ("04_learning_outcomes", "learning_outcome", "learning_outcome_id", "competency_level_id", ("outcome", "outcome_si"), ("outcome", "outcome_en_gloss")),
            ("05_content_points", "content_point", "content_point_id", "competency_level_id", ("content", "content_si"), ("content", "content_en_gloss")),
            ("06_lessons", "lesson", "lesson_id", "", ("lesson_title", "lesson_title_si"), ("lesson_title", "lesson_title_en_gloss")),
            ("07_activities", "source_activity", "activity_id", "lesson_id", ("activity_title", "activity_title_si", "instructions_summary", "instructions_summary_si"), ("activity_title",)),
        ]
        for sheet, node_type, id_field, parent_field, primary_fields, english_fields in node_specs:
            for row in subject_raw[sheet]:
                node_id = clean(row.get(id_field))
                if not node_id:
                    continue
                parent = clean(row.get(parent_field)) if parent_field else ""
                parent_basis = "stated"
                if not parent and node_type in ("learning_outcome", "content_point"):
                    parent = derive_parent_from_id(node_id, node_type)
                    parent_basis = "inferred_from_id"
                refs = first(row, "competency_level_ids")
                text_primary = first(row, *primary_fields)
                text_en = first(row, *english_fields)
                if subject == "ENG":
                    text_en = text_primary
                nodes.append({
                    "curriculum_node_id": node_id,
                    "curriculum_version": CURRICULUM_VERSION,
                    "subject": subject,
                    "subject_name": SUBJECT_NAMES[subject],
                    "grade": 5,
                    "node_type": node_type,
                    "parent_node_id": parent,
                    "parent_basis": parent_basis if parent else "none",
                    "curriculum_refs": refs,
                    "official_ref": first(row, "official_ref"),
                    "text_primary": text_primary,
                    "text_en": text_en,
                    "content_type": first(row, "content_type"),
                    "source_id": first(row, "source_id"),
                    "source_page_from": first(row, "page_from"),
                    "source_page_to": first(row, "page_to"),
                    "source_page_basis": first(row, "page_basis"),
                    "allocated_periods": first(row, "allocated_periods"),
                    "mapping_basis": first(row, "link_basis"),
                    "materials_primary": first(
                        row, "materials_mentioned", "materials_mentioned_si"
                    ),
                    "extraction_confidence": source_confidence(row),
                    "extraction_notes": first(row, "extraction_notes"),
                    "trace_status": first(row, "trace_status") or "draft",
                    "verification_status": "needs_review",
                    "reviewed_by": first(row, "verified_by"),
                    "reviewed_on": first(row, "verified_on"),
                    "source_file": path.name,
                    "source_sheet": sheet,
                    "source_row_number": row["_source_row_number"],
                    "source_record": source_record(row),
                })

    return nodes, raw


GAME_DEFINITIONS = {
    "MAT": [
        ("Place Value Forge", "place_value", "Numbers - 1", "Build and decompose whole numbers."),
        ("Number Order Dash", "number_order", "Numbers - 2", "Compare and order whole numbers."),
        ("Addition Quest", "addition", "Addition", "Add whole numbers accurately."),
        ("Subtraction Rescue", "subtraction", "Subtraction", "Subtract with regrouping."),
        ("Multiplication Arena", "multiplication", "Multiplication", "Recall and apply multiplication facts."),
        ("Division Drop", "division", "Division", "Use division facts with exact quotients."),
        ("Fraction Builder", "fractions", "Fractions", "Represent fractions of sets and wholes."),
        ("Decimal Market", "decimals", "Decimals", "Recognize and compare tenths."),
        ("Time Keeper", "time", "Time", "Read time and calculate elapsed time."),
        ("Money Shop", "money", "Money", "Calculate totals and change."),
        ("Measure Master", "measurement", "Length", "Convert and compare common measurements."),
        ("Pattern and Shape Lab", "patterns_shapes", "Number patterns", "Continue patterns and identify shape properties."),
    ],
    "ENG": [
        ("Sound Detective", "phonics", "sound", "Recognize sounds in familiar words."),
        ("Word Builder", "word_builder", "word", "Construct familiar words from letters."),
        ("Vocabulary Match", "vocabulary", "familiar", "Connect common words with meanings."),
        ("Sentence Scramble", "sentence_order", "sentence", "Arrange words into meaningful sentences."),
        ("Grammar Sort", "grammar_sort", "grammar", "Classify common word types."),
        ("Missing Word", "missing_word", "sentence", "Select a word that completes a sentence."),
        ("Punctuation Patrol", "punctuation", "punctuation", "Use capitals and end punctuation."),
        ("Spelling Sprint", "spelling", "word", "Recognize the correct spelling of familiar words."),
        ("Reading Clue Hunt", "reading", "read", "Read a short text and answer a question."),
        ("Dialogue Builder", "dialogue", "speak", "Choose an appropriate conversational response."),
        ("Listening Choice", "listening", "listen", "Listen for key information."),
        ("Topic Adventure", "topic_vocabulary", "topic", "Use vocabulary from Grade 5 themes."),
    ],
    "SIN": [
        ("සවන් දීමේ අභියෝගය", "activity_clue", "Listening", "Identify an activity from a listening or speaking clue."),
        ("කතන්දර මඟ", "lesson_activity", "story", "Connect lesson themes with learning activities."),
        ("වචන යුගල", "term_match", "Vocabulary", "Match Sinhala curriculum terms with their glosses."),
        ("කියවීමේ සලකුණු", "outcome_match", "Reading", "Connect a learning outcome with its description."),
        ("ක්‍රියාකාරකම හඳුනාගන්න", "activity_instruction", "activity", "Identify activities from their instructions."),
        ("පාඩම් සිතියම", "activity_lesson", "lesson", "Classify an activity under its lesson."),
        ("අනුපිළිවෙළ ගමන", "activity_sequence", "sequence", "Place lesson activities in source order."),
        ("සන්නිවේදන තේරීම", "communication", "communication", "Select a curriculum-aligned communication activity."),
        ("භාෂා කුසලතා", "skill_sort", "language", "Sort activities by language-skill lesson."),
        ("රසවින්දන සෙවීම", "appreciation", "appreciative", "Recognize appreciative language activities."),
        ("නිර්මාණාත්මක භාෂාව", "creative_language", "creative", "Recognize creative language-use activities."),
        ("සිංහල පුනරීක්ෂණය", "mixed_review", "review", "Review lesson, activity and outcome relationships."),
    ],
    "ENV": [
        ("යහපැවැත්ම මඟ", "activity_lesson", "Our Well-being", "Classify well-being activities."),
        ("අපි ලක් වැසියෝ", "activity_lesson", "We Sri Lankans", "Explore Sri Lankan identity and heritage."),
        ("පැළෑටි හා ගස්වැල්", "activity_lesson", "Plants and Trees", "Classify plant-related activities."),
        ("අප අවට සත්තු", "activity_lesson", "Animals Around Us", "Classify animal-related activities."),
        ("අප පරිහරණය කරන දේ", "activity_lesson", "Things We Use", "Recognize materials and everyday objects."),
        ("අපේ ආහාර", "activity_lesson", "Our Food", "Recognize food and nutrition activities."),
        ("ආරක්ෂාව සහ රැකවරණය", "activity_lesson", "Safety and Care", "Recognize safe actions and care."),
        ("ජල අභියෝගය", "activity_lesson", "Water", "Explore uses and care of water."),
        ("වැඩ පහසු කරමු", "activity_lesson", "Making Work Easier", "Recognize tools and simple methods."),
        ("ගෙවත්ත ගොඩනඟමු", "activity_lesson", "Home Garden", "Explore home-garden activities."),
        ("පෘථිවිය හා අහස", "activity_lesson", "Earth and Sky", "Classify Earth-and-sky observations."),
        ("පරිසර පුනරීක්ෂණය", "mixed_review", "review", "Review environmental themes and activities."),
    ],
}


def find_target_nodes(nodes_by_subject, subject, keyword):
    keyword_l = keyword.lower()
    candidates = []
    for node in nodes_by_subject[subject]:
        haystack = " ".join([node["text_primary"], node["text_en"], node["node_type"]]).lower()
        if keyword_l in haystack:
            candidates.append(node["curriculum_node_id"])
    if not candidates:
        candidates = [n["curriculum_node_id"] for n in nodes_by_subject[subject] if n["node_type"] in ("competency_level", "learning_outcome")]
    return candidates[:8]


def build_game_catalog(nodes):
    by_subject = defaultdict(list)
    for node in nodes:
        by_subject[node["subject"]].append(node)
    games = []
    for subject in SUBJECTS:
        for index, (title, mechanic, keyword, objective) in enumerate(GAME_DEFINITIONS[subject], start=1):
            refs = find_target_nodes(by_subject, subject, keyword)
            games.append({
                "game_id": f"G5-{subject}-G{index:02d}",
                "game_version": PACKAGE_VERSION,
                "curriculum_version": CURRICULUM_VERSION,
                "skill_taxonomy_version": SKILL_TAXONOMY_VERSION,
                "subject": subject,
                "subject_name": SUBJECT_NAMES[subject],
                "grade": 5,
                "title": title,
                "mechanic": mechanic,
                "learning_objective": objective,
                "target_curriculum_ids": ",".join(refs),
                "session_minutes": 8,
                "minimum_items": 20,
                "starter_items": 40,
                "difficulty_start": 1,
                "difficulty_end": 4,
                "scoring_rule": "10 points correct; 5 after one hint; streak bonus every 3 correct",
                "hint_rule": "One hint per item; remove one distractor or reveal a structural clue",
                "required_modalities": "text,audio" if mechanic in ("listening", "phonics", "activity_clue") else "text",
                "language": LANGUAGES[subject],
                "content_status": "prototype",
                "publication_status": "needs_review",
            })
    return games


def base_item(game, number, refs, lesson_id="", source_activity_id=""):
    return {
        "item_id": f"{game['game_id']}-I{number:03d}",
        "item_version": PACKAGE_VERSION,
        "curriculum_version": CURRICULUM_VERSION,
        "skill_taxonomy_version": SKILL_TAXONOMY_VERSION,
        "game_id": game["game_id"],
        "subject": game["subject"],
        "language": game["language"],
        "lesson_id": lesson_id,
        "curriculum_refs": ",".join(refs),
        "source_activity_id": source_activity_id,
        "mechanic": game["mechanic"],
        "prompt_primary": "",
        "prompt_en": "",
        "stimulus_primary": "",
        "stimulus_en": "",
        "options_json": "[]",
        "correct_answer_json": "[]",
        "feedback_correct": "Correct.",
        "feedback_incorrect": "Try again and use the hint.",
        "hint": "",
        "difficulty": 1 + ((number - 1) // 10),
        "cognitive_level": "understand" if number <= 20 else "apply",
        "estimated_seconds": 30,
        "generation_method": "rule_based",
        "source_confidence": "high",
        "answer_verified": False,
        "language_reviewed": False,
        "asset_status": "not_required",
        "publication_status": "needs_review",
        "item_role": "practice_candidate",
    }


def math_refs(game):
    refs = [r for r in game["target_curriculum_ids"].split(",") if r]
    return refs[:3]


def make_math_items(game):
    items = []
    refs = math_refs(game)
    mechanic = game["mechanic"]
    for i in range(1, 41):
        item = base_item(game, i, refs)
        item.update({"language": "bilingual", "answer_verified": True, "source_confidence": "high"})
        difficulty = item["difficulty"]
        if mechanic == "place_value":
            number = RNG.randint(1000, 9999 if difficulty < 3 else 99999)
            digits = list(str(number))
            position = RNG.randrange(len(digits))
            digit = int(digits[position])
            value = digit * (10 ** (len(digits) - position - 1))
            options = make_options(str(value), [str(digit), str(value * 10), str(max(1, value // 10)), str(number)])
            item.update(prompt_primary=f"{number} සංඛ්‍යාවේ {digit} ඉලක්කමේ ස්ථානීය අගය කුමක් ද?", prompt_en=f"What is the place value of digit {digit} in {number}?", stimulus_primary=str(number), stimulus_en=str(number), options_json=json.dumps(options, ensure_ascii=False), correct_answer_json=json.dumps([str(value)]), hint="Count places from the right: ones, tens, hundreds, thousands.")
        elif mechanic == "number_order":
            values = RNG.sample(range(100, 9999 if difficulty < 3 else 99999), 4)
            correct = " < ".join(map(str, sorted(values)))
            variants = []
            while len(variants) < 4:
                order = values[:]
                RNG.shuffle(order)
                candidate = " < ".join(map(str, order))
                if candidate != correct and candidate not in variants:
                    variants.append(candidate)
            item.update(prompt_primary="සංඛ්‍යා කුඩාම සිට විශාලම දක්වා සකසන්න.", prompt_en="Arrange the numbers from smallest to largest.", stimulus_primary=", ".join(map(str, values)), stimulus_en=", ".join(map(str, values)), options_json=json.dumps(make_options(correct, variants), ensure_ascii=False), correct_answer_json=json.dumps([correct]), hint="Compare the highest place value first.")
        elif mechanic == "addition":
            limit = 999 if difficulty < 3 else 9999
            a, b = RNG.randint(20, limit), RNG.randint(20, limit)
            answer = a + b
            options = make_options(str(answer), [str(answer + 10), str(answer - 10), str(answer + 1), str(abs(a - b))])
            item.update(prompt_primary=f"{a} + {b} හි අගය සොයන්න.", prompt_en=f"Calculate {a} + {b}.", stimulus_primary=f"{a} + {b}", stimulus_en=f"{a} + {b}", options_json=json.dumps(options), correct_answer_json=json.dumps([str(answer)]), hint="Add ones first, then tens, hundreds and thousands.")
        elif mechanic == "subtraction":
            a = RNG.randint(100, 999 if difficulty < 3 else 9999)
            b = RNG.randint(10, a)
            answer = a - b
            options = make_options(str(answer), [str(answer + 10), str(max(0, answer - 10)), str(a + b), str(answer + 1)])
            item.update(prompt_primary=f"{a} - {b} හි අගය සොයන්න.", prompt_en=f"Calculate {a} - {b}.", stimulus_primary=f"{a} - {b}", stimulus_en=f"{a} - {b}", options_json=json.dumps(options), correct_answer_json=json.dumps([str(answer)]), hint="Regroup when the upper digit is smaller.")
        elif mechanic == "multiplication":
            a = RNG.randint(2, 12 if difficulty < 3 else 25)
            b = RNG.randint(2, 12 if difficulty < 4 else 20)
            answer = a * b
            options = make_options(str(answer), [str(answer + a), str(max(0, answer - a)), str(a + b), str(answer + b)])
            item.update(prompt_primary=f"{a} × {b} හි අගය සොයන්න.", prompt_en=f"Calculate {a} × {b}.", stimulus_primary=f"{a} × {b}", stimulus_en=f"{a} × {b}", options_json=json.dumps(options), correct_answer_json=json.dumps([str(answer)]), hint=f"Think of {b} groups of {a}.")
        elif mechanic == "division":
            divisor = RNG.randint(2, 12)
            quotient = RNG.randint(2, 12 if difficulty < 3 else 30)
            dividend = divisor * quotient
            options = make_options(str(quotient), [str(divisor), str(quotient + 1), str(max(1, quotient - 1)), str(dividend - divisor)])
            item.update(prompt_primary=f"{dividend} ÷ {divisor} හි අගය සොයන්න.", prompt_en=f"Calculate {dividend} ÷ {divisor}.", stimulus_primary=f"{dividend} ÷ {divisor}", stimulus_en=f"{dividend} ÷ {divisor}", options_json=json.dumps(options), correct_answer_json=json.dumps([str(quotient)]), hint=f"What number multiplied by {divisor} gives {dividend}?")
        elif mechanic == "fractions":
            denominator = RNG.choice([2, 4, 10])
            groups = RNG.randint(1, 5)
            total = denominator * groups
            numerator = RNG.randint(1, denominator - 1)
            answer = total * numerator // denominator
            options = make_options(str(answer), [str(total // denominator), str(total - answer), str(answer + 1), str(max(1, answer - 1))])
            item.update(prompt_primary=f"වස්තු {total} කින් {numerator}/{denominator} කීය ද?", prompt_en=f"What is {numerator}/{denominator} of {total}?", stimulus_primary=f"{numerator}/{denominator} of {total}", stimulus_en=f"{numerator}/{denominator} of {total}", options_json=json.dumps(options), correct_answer_json=json.dumps([str(answer)]), hint=f"Divide {total} into {denominator} equal groups, then take {numerator} group(s).")
        elif mechanic == "decimals":
            a, b = RNG.randint(0, 9), RNG.randint(0, 9)
            if i % 2:
                correct = str(round((a + b) / 10, 1))
                prompt = f"0.{a} + 0.{b}"
            else:
                correct = f"0.{max(a,b)}"
                prompt = f"Which is greater: 0.{a} or 0.{b}?"
            options = make_options(correct, [f"0.{min(a,b)}", str(a + b), f"{a}.{b}", "1.0"])
            item.update(prompt_primary=prompt, prompt_en=prompt, stimulus_primary=prompt, stimulus_en=prompt, options_json=json.dumps(options), correct_answer_json=json.dumps([correct]), hint="Tenths are the first digit after the decimal point.")
        elif mechanic == "time":
            start_hour = RNG.randint(7, 15)
            start_min = RNG.choice([0, 15, 30, 45])
            elapsed = RNG.choice([15, 30, 45, 60, 90, 120])
            total = start_hour * 60 + start_min + elapsed
            answer = f"{(total // 60) % 24:02d}:{total % 60:02d}"
            start = f"{start_hour:02d}:{start_min:02d}"
            options = make_options(answer, [f"{(total // 60 + 1) % 24:02d}:{total % 60:02d}", f"{(total // 60) % 24:02d}:{(total + 15) % 60:02d}", start])
            item.update(prompt_primary=f"{start} සිට මිනිත්තු {elapsed} කට පසු වේලාව කුමක් ද?", prompt_en=f"What time is {elapsed} minutes after {start}?", stimulus_primary=start, stimulus_en=start, options_json=json.dumps(options), correct_answer_json=json.dumps([answer]), hint="Add the minutes first; exchange 60 minutes for one hour.")
        elif mechanic == "money":
            price1 = RNG.randrange(25, 501, 5)
            price2 = RNG.randrange(20, 301, 5)
            total = price1 + price2
            if i % 2:
                answer = total
                prompt_en = f"An item costs Rs. {price1} and another costs Rs. {price2}. What is the total?"
            else:
                paid = ((total + 99) // 100) * 100
                answer = paid - total
                prompt_en = f"The total is Rs. {total}. You pay Rs. {paid}. What is the change?"
            options = make_options(str(answer), [str(answer + 10), str(max(0, answer - 10)), str(total), str(price1)])
            item.update(prompt_primary=prompt_en, prompt_en=prompt_en, stimulus_primary=f"Rs. {price1}, Rs. {price2}", stimulus_en=f"Rs. {price1}, Rs. {price2}", options_json=json.dumps(options), correct_answer_json=json.dumps([str(answer)]), hint="For change, subtract the total cost from the amount paid.")
        elif mechanic == "measurement":
            kind = i % 3
            n = RNG.randint(2, 20)
            if kind == 0:
                prompt_en, answer = f"Convert {n} metres to centimetres.", n * 100
                unit = "cm"
            elif kind == 1:
                prompt_en, answer = f"Convert {n} kilograms to grams.", n * 1000
                unit = "g"
            else:
                prompt_en, answer = f"Convert {n} litres to millilitres.", n * 1000
                unit = "ml"
            correct = f"{answer} {unit}"
            options = make_options(correct, [f"{n * 10} {unit}", f"{n * 100} {unit}", f"{n} {unit}"])
            item.update(prompt_primary=prompt_en, prompt_en=prompt_en, stimulus_primary=str(n), stimulus_en=str(n), options_json=json.dumps(options), correct_answer_json=json.dumps([correct]), hint="Use 1 m = 100 cm and 1 kg or L = 1000 smaller units.")
        else:
            start = RNG.randint(1, 30)
            step = RNG.randint(2, 10)
            seq = [start + step * k for k in range(4)]
            answer = start + step * 4
            options = make_options(str(answer), [str(answer + step), str(answer - step), str(answer + 1), str(answer * 2)])
            item.update(prompt_primary="රටාවේ ඊළඟ සංඛ්‍යාව සොයන්න.", prompt_en="Find the next number in the pattern.", stimulus_primary=", ".join(map(str, seq)) + ", …", stimulus_en=", ".join(map(str, seq)) + ", …", options_json=json.dumps(options), correct_answer_json=json.dumps([str(answer)]), hint="Find the common difference between neighbouring numbers.")
        items.append(item)
    return items


ENGLISH_WORDS = [
    ("island", "land surrounded by water"), ("forest", "a large area covered with trees"),
    ("kind", "helpful and caring"), ("share", "give part of something to another"),
    ("festival", "a special time of celebration"), ("athlete", "a person who takes part in sport"),
    ("healthy", "well and strong"), ("market", "a place where things are bought and sold"),
    ("helmet", "protective covering worn on the head"), ("journey", "travel from one place to another"),
    ("future", "the time that has not happened yet"), ("river", "a natural stream of water"),
    ("friend", "a person you like and trust"), ("team", "a group working or playing together"),
    ("doctor", "a person who treats illness"), ("teacher", "a person who helps others learn"),
    ("careful", "taking care to avoid danger"), ("country", "a nation with its own land"),
    ("garden", "a place where plants are grown"), ("animal", "a living creature that can move"),
    ("honest", "truthful and fair"), ("brave", "ready to face danger"),
    ("quiet", "making little or no sound"), ("rapid", "very fast"),
    ("protect", "keep safe from harm"), ("celebrate", "do something special for an occasion"),
    ("exercise", "activity that keeps the body fit"), ("traffic", "vehicles moving on a road"),
    ("planet", "a large body moving around a star"), ("invent", "create something new"),
    ("village", "a small settlement"), ("ocean", "a very large area of salt water"),
    ("respect", "show care for someone or something"), ("recycle", "use material again"),
    ("balance", "remain steady without falling"), ("energy", "power needed for activity"),
    ("visitor", "a person who comes to a place"), ("message", "information sent to someone"),
    ("shelter", "a place that gives protection"), ("discover", "find something for the first time"),
]


ENGLISH_SENTENCES = [
    "Sri Lanka is our beautiful island",
    "Trees give shade to people and animals",
    "We should care for our friends",
    "The children celebrate the festival together",
    "Our team practises every afternoon",
    "Clean food helps us stay healthy",
    "The family visits the fair on Saturday",
    "Wear a helmet when you ride a bicycle",
    "They travelled around the world",
    "We will build a kinder future",
]


def make_english_items(game):
    items = []
    refs = [r for r in game["target_curriculum_ids"].split(",") if r][:3]
    words = [w for w, _ in ENGLISH_WORDS]
    for i in range(1, 41):
        item = base_item(game, i, refs)
        word, definition = ENGLISH_WORDS[i - 1]
        mechanic = game["mechanic"]
        if mechanic == "phonics":
            target = word[0]
            choices = make_options(word, [w for w in words if not w.startswith(target)])
            item.update(prompt_primary=f"Which word begins with the sound /{target}/?", prompt_en=f"Which word begins with the sound /{target}/?", stimulus_primary=target, stimulus_en=target, options_json=json.dumps(choices), correct_answer_json=json.dumps([word]), hint=f"Say /{target}/ slowly, then say each option.")
        elif mechanic == "word_builder":
            letters = list(word)
            RNG.shuffle(letters)
            item.update(prompt_primary="Build a word from these letters.", prompt_en="Build a word from these letters.", stimulus_primary=" ".join(letters), stimulus_en=" ".join(letters), options_json=json.dumps(make_options(word, words)), correct_answer_json=json.dumps([word]), hint=definition)
        elif mechanic == "vocabulary":
            item.update(prompt_primary=f"Which word means “{definition}”?", prompt_en=f"Which word means “{definition}”?", stimulus_primary=definition, stimulus_en=definition, options_json=json.dumps(make_options(word, words)), correct_answer_json=json.dumps([word]), hint=f"The word begins with {word[0].upper()}.")
        elif mechanic == "sentence_order":
            sentence = ENGLISH_SENTENCES[(i - 1) % len(ENGLISH_SENTENCES)]
            words_in_sentence = sentence.split()
            shuffled = words_in_sentence[:]
            RNG.shuffle(shuffled)
            answer = sentence + "."
            item.update(prompt_primary="Arrange the words to make a sentence.", prompt_en="Arrange the words to make a sentence.", stimulus_primary=" / ".join(shuffled), stimulus_en=" / ".join(shuffled), options_json=json.dumps(shuffled), correct_answer_json=json.dumps([answer]), hint="Start with a capital letter and finish with a full stop.")
        elif mechanic == "grammar_sort":
            categories = ["noun", "verb", "adjective"]
            category = "adjective" if word in {"kind", "healthy", "careful", "honest", "brave", "quiet", "rapid"} else ("verb" if word in {"share", "protect", "celebrate", "exercise", "invent", "respect", "recycle", "balance", "discover"} else "noun")
            item.update(prompt_primary=f"What kind of word is “{word}”?", prompt_en=f"What kind of word is “{word}”?", stimulus_primary=word, stimulus_en=word, options_json=json.dumps(categories), correct_answer_json=json.dumps([category]), hint="A noun names; a verb shows action; an adjective describes.")
        elif mechanic == "missing_word":
            sentence = f"The word _____ means {definition}."
            item.update(prompt_primary="Choose the missing word.", prompt_en="Choose the missing word.", stimulus_primary=sentence, stimulus_en=sentence, options_json=json.dumps(make_options(word, words)), correct_answer_json=json.dumps([word]), hint=f"It begins with {word[0]}.")
        elif mechanic == "punctuation":
            base = ENGLISH_SENTENCES[(i - 1) % len(ENGLISH_SENTENCES)]
            correct = base + "."
            wrong = [base.lower(), base + "?", base.upper() + "."]
            item.update(prompt_primary="Choose the correctly written sentence.", prompt_en="Choose the correctly written sentence.", stimulus_primary=base.lower(), stimulus_en=base.lower(), options_json=json.dumps(make_options(correct, wrong)), correct_answer_json=json.dumps([correct]), hint="Check the capital letter and end punctuation.")
        elif mechanic == "spelling":
            variants = [word[:-1], word + word[-1], (word[1] + word[0] + word[2:]) if len(word) > 2 else word + "e"]
            item.update(prompt_primary="Choose the correctly spelled word.", prompt_en="Choose the correctly spelled word.", stimulus_primary=definition, stimulus_en=definition, options_json=json.dumps(make_options(word, variants)), correct_answer_json=json.dumps([word]), hint=f"The word has {len(word)} letters.")
        elif mechanic == "reading":
            topic = ENGLISH_SENTENCES[(i - 1) % len(ENGLISH_SENTENCES)] + "."
            target = topic.split()[0]
            item.update(prompt_primary="Read the sentence. Which word begins the sentence?", prompt_en="Read the sentence. Which word begins the sentence?", stimulus_primary=topic, stimulus_en=topic, options_json=json.dumps(make_options(target, words)), correct_answer_json=json.dumps([target]), hint="Look at the first word after the beginning of the line.")
        elif mechanic == "dialogue":
            prompts = [("How are you?", "I am well, thank you."), ("May I borrow your pencil?", "Yes, you may."), ("Thank you for helping me.", "You are welcome."), ("Where is the library?", "It is next to the office."), ("Would you like to play?", "Yes, I would.")]
            question, answer = prompts[(i - 1) % len(prompts)]
            distractors = [p[1] for p in prompts if p[1] != answer]
            item.update(prompt_primary="Choose the best reply.", prompt_en="Choose the best reply.", stimulus_primary=question, stimulus_en=question, options_json=json.dumps(make_options(answer, distractors)), correct_answer_json=json.dumps([answer]), hint="Choose the reply that directly answers the speaker.")
        elif mechanic == "listening":
            sentence = ENGLISH_SENTENCES[(i - 1) % len(ENGLISH_SENTENCES)] + "."
            target = sentence.split()[RNG.randrange(len(sentence.split()))].strip(".")
            item.update(prompt_primary=f"Listen and choose the word you heard: {target}", prompt_en=f"Listen and choose the word you heard: {target}", stimulus_primary=sentence, stimulus_en=sentence, options_json=json.dumps(make_options(target, words)), correct_answer_json=json.dumps([target]), hint="Replay the audio once and listen for the target word.", asset_status="audio_required")
        else:
            lesson_topics = ["Sri Lanka", "nature", "sharing", "celebration", "sports", "health", "fair", "safety", "world", "future"]
            topic = lesson_topics[(i - 1) % len(lesson_topics)]
            item.update(prompt_primary=f"Choose the word that best fits the topic “{topic}”.", prompt_en=f"Choose the word that best fits the topic “{topic}”.", stimulus_primary=topic, stimulus_en=topic, options_json=json.dumps(make_options(word, words)), correct_answer_json=json.dumps([word]), hint=definition)
        item.update(answer_verified=True, language_reviewed=True, source_confidence="high")
        items.append(item)
    return items


def index_raw(raw, subject):
    lessons = raw[subject]["06_lessons"]
    activities = raw[subject]["07_activities"]
    outcomes = raw[subject]["04_learning_outcomes"]
    content = raw[subject]["05_content_points"]
    terms = raw[subject]["08_terms"]
    lesson_title = {}
    for row in lessons:
        lesson_title[first(row, "lesson_id")] = {
            "primary": first(row, "lesson_title", "lesson_title_si"),
            "en": first(row, "lesson_title", "lesson_title_en_gloss"),
        }
    return lessons, activities, outcomes, content, terms, lesson_title


def make_source_derived_items(game, raw):
    subject = game["subject"]
    lessons, activities, outcomes, content, terms, lesson_titles = index_raw(raw, subject)
    valid_activities = [a for a in activities if first(a, "lesson_id") and first(a, "activity_title", "activity_title_si", "instructions_summary", "instructions_summary_si")]
    if not valid_activities:
        valid_activities = activities
    title_pool = [v["primary"] for v in lesson_titles.values() if v["primary"]]
    activity_title_pool = [first(a, "activity_title", "activity_title_si", "instructions_summary", "instructions_summary_si") for a in valid_activities]
    term_pool_en = [first(t, "term_en_gloss", "term") for t in terms if first(t, "term_en_gloss", "term")]
    items = []
    for i in range(1, 41):
        activity = valid_activities[(i - 1) % len(valid_activities)]
        lesson_id = first(activity, "lesson_id")
        lesson = lesson_titles.get(lesson_id, {"primary": lesson_id, "en": lesson_id})
        activity_title = first(activity, "activity_title", "activity_title_si", "instructions_summary", "instructions_summary_si")
        instruction = first(activity, "instructions_summary", "instructions_summary_si") or activity_title
        refs = [r.strip() for r in first(activity, "competency_level_ids").split(",") if r.strip()]
        if not refs:
            refs = [lesson_id] if lesson_id else [r for r in game["target_curriculum_ids"].split(",") if r][:1]
        item = base_item(game, i, refs, lesson_id, first(activity, "activity_id"))
        item.update(source_confidence=source_confidence(activity), generation_method="curriculum_adaptation")
        mechanic = game["mechanic"]
        if mechanic == "term_match" and terms:
            term = terms[(i - 1) % len(terms)]
            primary = first(term, "term", "term_si")
            english = first(term, "term_en_gloss", "term")
            item.update(prompt_primary=f"“{primary}” යන්නට ගැළපෙන අර්ථය තෝරන්න.", prompt_en=f"Choose the matching meaning for “{primary}”.", stimulus_primary=primary, stimulus_en=english, options_json=json.dumps(make_options(english, term_pool_en), ensure_ascii=False), correct_answer_json=json.dumps([english], ensure_ascii=False), hint="Compare the key idea in each option.")
        elif mechanic == "activity_instruction":
            item.update(prompt_primary="උපදෙස්වලට ගැළපෙන ක්‍රියාකාරකම තෝරන්න.", prompt_en="Choose the activity that matches the instruction.", stimulus_primary=instruction, stimulus_en="", options_json=json.dumps(make_options(activity_title, activity_title_pool), ensure_ascii=False), correct_answer_json=json.dumps([activity_title], ensure_ascii=False), hint="Look for the action described in the instruction.")
        elif mechanic == "activity_sequence":
            same = [a for a in valid_activities if first(a, "lesson_id") == lesson_id][:4]
            ordered = [first(a, "activity_title", "activity_title_si", "instructions_summary", "instructions_summary_si") for a in same if first(a, "activity_title", "activity_title_si", "instructions_summary", "instructions_summary_si")]
            if len(ordered) < 3:
                ordered = activity_title_pool[:4]
            shuffled = ordered[:]
            RNG.shuffle(shuffled)
            item.update(prompt_primary="ක්‍රියාකාරකම් මූලාශ්‍ර අනුපිළිවෙළට සකසන්න.", prompt_en="Arrange the activities in source order.", stimulus_primary=" | ".join(shuffled), stimulus_en="", options_json=json.dumps(shuffled, ensure_ascii=False), correct_answer_json=json.dumps(ordered, ensure_ascii=False), hint="Use the lesson progression from introduction to practice.")
        elif mechanic == "outcome_match" and outcomes:
            outcome = outcomes[(i - 1) % len(outcomes)]
            primary = first(outcome, "outcome", "outcome_si")
            english = first(outcome, "outcome_en_gloss", "outcome")
            en_pool = [first(o, "outcome_en_gloss", "outcome") for o in outcomes if first(o, "outcome_en_gloss", "outcome")]
            item.update(prompt_primary="මෙම ඉගෙනුම් ප්‍රතිඵලයට ගැළපෙන විස්තරය තෝරන්න.", prompt_en="Choose the matching description for this learning outcome.", stimulus_primary=primary, stimulus_en=english, options_json=json.dumps(make_options(english or primary, en_pool or [primary]), ensure_ascii=False), correct_answer_json=json.dumps([english or primary], ensure_ascii=False), hint="Identify the main action expected from the learner.")
            item["curriculum_refs"] = first(outcome, "learning_outcome_id")
        else:
            correct = lesson["primary"] or lesson_id
            item.update(prompt_primary="මෙම ක්‍රියාකාරකමට ගැළපෙන පාඩම් මාතෘකාව තෝරන්න.", prompt_en="Choose the lesson topic that best matches this activity.", stimulus_primary=instruction, stimulus_en="", options_json=json.dumps(make_options(correct, title_pool), ensure_ascii=False), correct_answer_json=json.dumps([correct], ensure_ascii=False), hint="Focus on the main subject of the activity.")
        item.update(answer_verified=False, language_reviewed=False, publication_status="needs_review")
        items.append(item)
    return items


def build_items(games, raw):
    items = []
    for game in games:
        if game["subject"] == "MAT":
            items.extend(make_math_items(game))
        elif game["subject"] == "ENG":
            items.extend(make_english_items(game))
        else:
            items.extend(make_source_derived_items(game, raw))
    return items


def snake_case(value):
    return re.sub(r"[^a-z0-9]+", "_", clean(value).lower()).strip("_")


def split_ids(value):
    return [part.strip() for part in clean(value).split(",") if part.strip()]


def read_master_table(workbook, sheet_name, header_row=5):
    ws = workbook[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    headers = [snake_case(value) for value in rows[header_row - 1]]
    records = []
    for row_number, row in enumerate(rows[header_row:], start=header_row + 1):
        if not any(value is not None and clean(value) for value in row):
            continue
        record = {
            header: json_value(value)
            for header, value in zip(headers, row)
            if header
        }
        record["source_sheet"] = sheet_name
        record["source_row_number"] = row_number
        records.append(record)
    return records


def read_master_annotations(workbook):
    annotations = []
    for sheet_name in ("Overview", "ReadMe"):
        ws = workbook[sheet_name]
        for row_number, row in enumerate(ws.iter_rows(values_only=True), start=1):
            values = [json_value(value) for value in row]
            if not any(value != "" for value in values):
                continue
            annotations.append(
                {
                    "annotation_id": f"MASTER-{snake_case(sheet_name).upper()}-R{row_number:04d}",
                    "source_workbook": MASTER_WORKBOOK_NAME,
                    "source_sheet": sheet_name,
                    "source_row_number": row_number,
                    "values": values,
                    "content_role": "source_reference_text",
                    "instruction_authority": "none",
                }
            )
    return annotations


def collect_master_workbook(source_dir):
    from openpyxl import load_workbook

    path = source_dir / MASTER_WORKBOOK_NAME
    workbook = load_workbook(path, read_only=True, data_only=True)
    traceability = read_master_table(workbook, "Traceability")
    qa_issues = read_master_table(workbook, "QA Queue")
    asset_requirements = read_master_table(workbook, "Assets")

    item_curriculum_mappings = []
    for row in traceability:
        item_curriculum_mappings.append(
            {
                "item_curriculum_mapping_id": row["trace_id"],
                "item_id": row["item_id"],
                "game_id": row["game_id"],
                "subject": row["subject"],
                "curriculum_node_id": row["curriculum_node_id"],
                "relationship": row["relationship"],
                "mapping_method": row["mapping_method"],
                "mapping_confidence": row["mapping_confidence"],
                "review_status": row["review_status"],
                "mapping_version": CURRICULUM_VERSION,
                "source_workbook": MASTER_WORKBOOK_NAME,
                "source_sheet": row["source_sheet"],
                "source_row_number": row["source_row_number"],
            }
        )

    normalized_qa = []
    for row in qa_issues:
        normalized_qa.append(
            {
                "qa_issue_id": row["qa_id"],
                "scope": row["scope"],
                "subject": row["subject"],
                "game_id": row["game_id"],
                "item_id": row["item_id"],
                "issue_type": row["issue_type"],
                "severity": row["severity"],
                "description": row["description"],
                "suggested_action": row["suggested_action"],
                "status": row["status"],
                "source_ref": row["source_ref"],
                "source_workbook": MASTER_WORKBOOK_NAME,
                "source_sheet": row["source_sheet"],
                "source_row_number": row["source_row_number"],
            }
        )

    normalized_assets = []
    for row in asset_requirements:
        normalized_assets.append(
            {
                "asset_requirement_id": row["asset_requirement_id"],
                "game_id": row["game_id"],
                "subject": row["subject"],
                "asset_type": row["asset_type"],
                "required_count": row["required_count"],
                "specification": row["specification"],
                "language": row["language"],
                "license_requirement": row["license_requirement"],
                "status": row["status"],
                "source_workbook": MASTER_WORKBOOK_NAME,
                "source_sheet": row["source_sheet"],
                "source_row_number": row["source_row_number"],
            }
        )
    return (
        item_curriculum_mappings,
        normalized_qa,
        normalized_assets,
        read_master_annotations(workbook),
    )


def collect_source_entities(raw, source_dir):
    curriculum_sources = []
    curriculum_terms = []
    source_issues = []
    for subject in SUBJECTS:
        workbook_name = f"grade5_{subject}_syllabus_trace_v1.xlsx"
        workbook_path = source_dir / workbook_name
        workbook_sha256 = hashlib.sha256(workbook_path.read_bytes()).hexdigest()
        for row in raw[subject]["01_sources"]:
            curriculum_sources.append(
                {
                    "source_id": first(row, "source_id"),
                    "subject": subject,
                    "grade": 5,
                    "document_file_name": first(row, "file_name"),
                    "document_type": first(row, "document_type"),
                    "title_primary": first(row, "title", "title_si"),
                    "title_en": first(row, "title", "title_en_gloss"),
                    "authority": first(row, "authority"),
                    "edition_year": first(row, "edition_year"),
                    "language": first(row, "language"),
                    "pages_in_file": first(row, "pages_in_file"),
                    "pages_traced": first(row, "pages_traced"),
                    "text_quality": first(row, "text_quality"),
                    "extraction_method": first(row, "extraction_method"),
                    "notes": first(row, "notes"),
                    "source_uri": "",
                    "source_document_sha256": "",
                    "license_note": "",
                    "approval_status": "needs_review",
                    "trace_workbook": workbook_name,
                    "trace_workbook_sha256": workbook_sha256,
                    "source_sheet": "01_sources",
                    "source_row_number": row["_source_row_number"],
                    "source_record": source_record(row),
                }
            )
        for row in raw[subject]["08_terms"]:
            curriculum_terms.append(
                {
                    "term_id": first(row, "term_id"),
                    "subject": subject,
                    "grade": 5,
                    "term_primary": first(row, "term", "term_si"),
                    "term_en": first(row, "term_en_gloss", "term"),
                    "definition_primary": first(row, "definition", "definition_si"),
                    "first_skill_candidate_id": first(
                        row, "first_competency_level_id"
                    ),
                    "first_lesson_id": first(row, "first_lesson_id"),
                    "source_id": first(row, "source_id"),
                    "source_page_from": first(row, "page_from"),
                    "trace_status": first(row, "trace_status") or "draft",
                    "source_workbook": workbook_name,
                    "source_sheet": "08_terms",
                    "source_row_number": row["_source_row_number"],
                    "source_record": source_record(row),
                }
            )
        for row in raw[subject]["09_issues"]:
            source_issues.append(
                {
                    "qa_issue_id": first(row, "issue_id"),
                    "scope": "source",
                    "subject": subject,
                    "game_id": "",
                    "item_id": first(row, "row_id") if first(row, "sheet") == "Game Items" else "",
                    "issue_type": first(row, "issue_type"),
                    "severity": "high" if first(row, "issue_type") == "unreadable" else "medium",
                    "description": first(row, "description"),
                    "suggested_action": first(row, "suggested_check"),
                    "status": "open",
                    "source_ref": f"{first(row, 'source_id')} p.{first(row, 'page')}",
                    "source_workbook": workbook_name,
                    "source_sheet": "09_issues",
                    "source_row_number": row["_source_row_number"],
                    "source_record": source_record(row),
                }
            )
    return curriculum_sources, curriculum_terms, source_issues


def normalize_games_and_items(games, items):
    normalized_games = []
    for game in games:
        row = dict(game)
        row["target_curriculum_node_ids"] = split_ids(row.pop("target_curriculum_ids"))
        row["required_modalities"] = split_ids(row["required_modalities"])
        normalized_games.append(row)

    game_by_id = {game["game_id"]: game for game in normalized_games}
    ordered_mechanics = {"sentence_order", "activity_sequence"}
    normalized_items = []
    for item in items:
        row = dict(item)
        row["curriculum_node_ids"] = split_ids(row.pop("curriculum_refs"))
        row["options"] = json.loads(row.pop("options_json"))
        row["correct_answers"] = json.loads(row.pop("correct_answer_json"))
        row["response_format"] = (
            "ordered_response" if row["mechanic"] in ordered_mechanics else "single_select"
        )
        modalities = game_by_id[row["game_id"]]["required_modalities"]
        row["required_asset_types"] = [
            modality for modality in modalities if modality != "text"
        ]
        row["asset_status"] = (
            "audio_required" if "audio" in row["required_asset_types"] else "not_required"
        )
        row["primary_skill_candidate_id"] = ""
        row["structured_error_mapping_status"] = "not_authored"
        normalized_items.append(row)
    return normalized_games, normalized_items


def build_skill_entities(nodes):
    skill_node_types = {"competency", "competency_level", "learning_outcome"}
    skill_nodes = [node for node in nodes if node["node_type"] in skill_node_types]
    skill_ids = {node["curriculum_node_id"] for node in skill_nodes}
    skills = []
    relationships = []
    source_mappings = []
    for node in skill_nodes:
        skill_id = node["curriculum_node_id"]
        parent_id = node["parent_node_id"] if node["parent_node_id"] in skill_ids else ""
        if not parent_id and "-L00-O" in skill_id:
            inferred_parent = skill_id.split("-L00-O", 1)[0]
            if inferred_parent in skill_ids:
                parent_id = inferred_parent
        skills.append(
            {
                "skill_candidate_id": skill_id,
                "gamification_skill_id": "",
                "taxonomy_version": SKILL_TAXONOMY_VERSION,
                "subject": node["subject"],
                "grade": node["grade"],
                "skill_type": node["node_type"],
                "parent_skill_candidate_id": parent_id,
                "label_primary": node["text_primary"],
                "label_en": node["text_en"],
                "observable_success_primary": (
                    node["text_primary"] if node["node_type"] == "learning_outcome" else ""
                ),
                "observable_success_en": (
                    node["text_en"] if node["node_type"] == "learning_outcome" else ""
                ),
                "prerequisite_skill_candidate_ids": [],
                "misconception_codes": ["unknown/other"],
                "eligible_support_action_types": [],
                "source_curriculum_node_id": skill_id,
                "mapping_status": "needs_review",
                "approval_status": "draft",
                "approved_by": "",
                "approved_on": "",
            }
        )
        if parent_id:
            relationships.append(
                {
                    "skill_relationship_id": f"{parent_id}--contains--{skill_id}",
                    "from_skill_candidate_id": parent_id,
                    "to_skill_candidate_id": skill_id,
                    "relationship_type": "contains",
                    "relationship_basis": node["parent_basis"],
                    "taxonomy_version": SKILL_TAXONOMY_VERSION,
                    "review_status": "needs_review",
                }
            )
        source_mappings.append(
            {
                "skill_source_mapping_id": f"{skill_id}--defines--{skill_id}",
                "skill_candidate_id": skill_id,
                "curriculum_node_id": skill_id,
                "mapping_type": "defines",
                "mapping_confidence": node["extraction_confidence"],
                "mapping_status": "needs_review",
                "mapping_version": SKILL_TAXONOMY_VERSION,
                "source_id": node["source_id"],
                "source_page_from": node["source_page_from"],
                "source_page_to": node["source_page_to"],
            }
        )
    return skills, relationships, source_mappings


def resolve_skill_candidates(curriculum_node_id, node_by_id, skill_ids):
    node = node_by_id.get(curriculum_node_id)
    if node is None:
        return []
    if curriculum_node_id in skill_ids:
        return [curriculum_node_id]
    candidates = []
    for ref in split_ids(node.get("curriculum_refs", "")):
        if ref in skill_ids:
            candidates.append(ref)
    parent_id = node.get("parent_node_id", "")
    if parent_id in skill_ids:
        candidates.append(parent_id)
    if parent_id and parent_id in node_by_id and not candidates:
        candidates.extend(resolve_skill_candidates(parent_id, node_by_id, skill_ids))
    if not candidates and "-L00-K" in curriculum_node_id:
        competency_id = curriculum_node_id.split("-L00-K", 1)[0]
        if competency_id in skill_ids:
            candidates.append(competency_id)
    return list(dict.fromkeys(candidates))


def build_item_skill_mappings(items, item_curriculum_mappings, nodes, skills):
    node_by_id = {node["curriculum_node_id"]: node for node in nodes}
    skill_ids = {skill["skill_candidate_id"] for skill in skills}
    mappings = []
    seen = set()
    by_item = defaultdict(list)
    for trace in item_curriculum_mappings:
        candidate_ids = resolve_skill_candidates(
            trace["curriculum_node_id"], node_by_id, skill_ids
        )
        for candidate_id in candidate_ids:
            key = (trace["item_id"], candidate_id, trace["relationship"])
            if key in seen:
                continue
            seen.add(key)
            mapping = {
                "item_skill_mapping_id": f"{trace['item_id']}--{candidate_id}--{trace['relationship']}",
                "item_id": trace["item_id"],
                "game_id": trace["game_id"],
                "skill_candidate_id": candidate_id,
                "gamification_skill_id": "",
                "relationship": trace["relationship"],
                "mapping_method": trace["mapping_method"],
                "mapping_confidence": trace["mapping_confidence"],
                "mapping_status": "needs_review",
                "mapping_version": SKILL_TAXONOMY_VERSION,
            }
            mappings.append(mapping)
            by_item[trace["item_id"]].append(mapping)
    for item in items:
        candidates = sorted(
            by_item[item["item_id"]],
            key=lambda value: (
                value["relationship"] != "assesses",
                value["skill_candidate_id"],
            ),
        )
        if candidates:
            item["primary_skill_candidate_id"] = candidates[0]["skill_candidate_id"]
    return mappings


def actual_text_languages(text):
    if re.search(r"[\u0D80-\u0DFF]", text):
        return ["si"]
    return ["en"] if clean(text) else []


def intended_languages(language):
    if language == "bilingual":
        return ["si", "en"]
    return [language]


def build_support_entities(items, item_skill_mappings):
    skills_by_item = defaultdict(list)
    for mapping in item_skill_mappings:
        skills_by_item[mapping["item_id"]].append(mapping["skill_candidate_id"])
    support_actions = []
    content_assets = []
    for item in items:
        asset_id = f"{item['item_id']}--HINT"
        action_id = f"{item['item_id']}--TARGETED-HINT"
        languages = actual_text_languages(item["hint"])
        skill_ids = sorted(set(skills_by_item[item["item_id"]]))
        content_assets.append(
            {
                "content_asset_id": asset_id,
                "content_version": PACKAGE_VERSION,
                "asset_type": "text_hint",
                "source_item_id": item["item_id"],
                "game_id": item["game_id"],
                "skill_candidate_ids": skill_ids,
                "misconception_codes": ["unknown/other"],
                "languages": languages,
                "intended_languages": intended_languages(item["language"]),
                "content_primary": item["hint"],
                "content_en": item["hint"] if "en" in languages else "",
                "answer_leakage_status": "needs_review",
                "age_appropriateness_status": "needs_review",
                "language_review_status": (
                    "reviewed" if item["language_reviewed"] and "en" in languages else "needs_review"
                ),
                "publication_status": "needs_review",
                "approved_by": "",
                "approved_on": "",
            }
        )
        support_actions.append(
            {
                "action_id": action_id,
                "catalogue_version": SUPPORT_CATALOGUE_VERSION,
                "action_type": "targeted_hint",
                "content_asset_id": asset_id,
                "languages": languages,
                "intended_languages": intended_languages(item["language"]),
                "skill_candidate_ids": skill_ids,
                "gamification_skill_ids": [],
                "misconception_codes": ["unknown/other"],
                "approved": False,
                "approval_status": "needs_review",
                "reference_priority": 1,
            }
        )
    return support_actions, content_assets


def build_verification_pair_candidates(items):
    groups = defaultdict(list)
    for item in items:
        skill_id = item["primary_skill_candidate_id"]
        if not skill_id:
            continue
        key = (
            item["game_id"],
            skill_id,
            item["difficulty"],
            item["cognitive_level"],
            item["response_format"],
            item["language"],
        )
        groups[key].append(item)
    pairs = []
    for group_items in groups.values():
        if len(group_items) < 2:
            continue
        ordered = sorted(group_items, key=lambda item: item["item_id"])
        for index, baseline in enumerate(ordered):
            verification = ordered[(index + 1) % len(ordered)]
            pairs.append(
                {
                    "verification_pair_id": f"{baseline['item_id']}--{verification['item_id']}",
                    "equivalence_version": EQUIVALENCE_VERSION,
                    "baseline_item_id": baseline["item_id"],
                    "verification_item_id": verification["item_id"],
                    "skill_candidate_id": baseline["primary_skill_candidate_id"],
                    "primary_skill_match": True,
                    "cognitive_operation_match": True,
                    "difficulty_band_match": True,
                    "response_format_match": True,
                    "language_reviewed": bool(
                        baseline["language_reviewed"] and verification["language_reviewed"]
                    ),
                    "expected_time_match": (
                        baseline["estimated_seconds"] == verification["estimated_seconds"]
                    ),
                    "unseen_item_required": True,
                    "unseen_verification_item": None,
                    "answer_overlap_acceptable": None,
                    "approved": False,
                    "review_status": "needs_review",
                    "reviewed_by": "",
                    "reviewed_on": "",
                }
            )
    return pairs


def generated_qa_issues(items, item_skill_mappings, support_actions, verification_pairs):
    mapped_item_ids = {mapping["item_id"] for mapping in item_skill_mappings}
    unresolved = [item for item in items if item["item_id"] not in mapped_item_ids]
    untranslated = [item for item in items if not item["stimulus_en"]]
    intended_language_mismatch = [
        action
        for action in support_actions
        if not set(action["intended_languages"]) <= set(action["languages"])
    ]
    return [
        {
            "qa_issue_id": "G5-PKG-ISS001",
            "scope": "dataset",
            "subject": "ALL",
            "game_id": "",
            "item_id": "",
            "issue_type": "unresolved_skill_mapping",
            "severity": "high",
            "description": f"{len(unresolved)} items have no source-supported skill candidate mapping.",
            "suggested_action": "Review and approve item-to-skill mappings without inferring unsupported links.",
            "status": "open",
            "source_ref": "generated package validation",
            "affected_count": len(unresolved),
        },
        {
            "qa_issue_id": "G5-PKG-ISS002",
            "scope": "dataset",
            "subject": "ALL",
            "game_id": "",
            "item_id": "",
            "issue_type": "missing_translation",
            "severity": "medium",
            "description": f"{len(untranslated)} items have no English stimulus equivalent.",
            "suggested_action": "Add and review language-equivalent stimulus text where required.",
            "status": "open",
            "source_ref": "generated package validation",
            "affected_count": len(untranslated),
        },
        {
            "qa_issue_id": "G5-PKG-ISS003",
            "scope": "model",
            "subject": "ALL",
            "game_id": "",
            "item_id": "",
            "issue_type": "missing_structured_error_mapping",
            "severity": "high",
            "description": "No expert-approved structured error to misconception mappings exist.",
            "suggested_action": "Complete double review and adjudication before enabling misconception diagnosis.",
            "status": "open",
            "source_ref": "misconception model release gate",
            "affected_count": len(items),
        },
        {
            "qa_issue_id": "G5-PKG-ISS004",
            "scope": "model",
            "subject": "ALL",
            "game_id": "",
            "item_id": "",
            "issue_type": "support_language_mismatch",
            "severity": "high",
            "description": f"{len(intended_language_mismatch)} support actions lack all intended delivery languages.",
            "suggested_action": "Author and review localized support assets before policy eligibility.",
            "status": "open",
            "source_ref": "support policy release gate",
            "affected_count": len(intended_language_mismatch),
        },
        {
            "qa_issue_id": "G5-PKG-ISS005",
            "scope": "model",
            "subject": "ALL",
            "game_id": "",
            "item_id": "",
            "issue_type": "verification_pairs_unapproved",
            "severity": "high",
            "description": f"All {len(verification_pairs)} generated verification-pair candidates require equivalence review.",
            "suggested_action": "Review skill, cognitive operation, format, language, time and answer overlap before approval.",
            "status": "open",
            "source_ref": "Independence Gain release gate",
            "affected_count": len(verification_pairs),
        },
    ]


FILE_ENTITY_SPECS = {
    "curriculum_sources.jsonl": (
        "curriculum_source",
        "source_id",
        ["source_id", "subject", "grade", "document_file_name", "approval_status"],
    ),
    "curriculum_nodes.jsonl": (
        "curriculum_node",
        "curriculum_node_id",
        [
            "curriculum_node_id",
            "curriculum_version",
            "subject",
            "grade",
            "node_type",
            "text_primary",
            "source_id",
            "verification_status",
        ],
    ),
    "curriculum_terms.jsonl": (
        "curriculum_term",
        "term_id",
        ["term_id", "subject", "term_primary", "source_id", "trace_status"],
    ),
    "skill_candidates.jsonl": (
        "skill_candidate",
        "skill_candidate_id",
        [
            "skill_candidate_id",
            "taxonomy_version",
            "subject",
            "grade",
            "skill_type",
            "label_primary",
            "mapping_status",
            "approval_status",
        ],
    ),
    "skill_relationships.jsonl": (
        "skill_relationship",
        "skill_relationship_id",
        [
            "skill_relationship_id",
            "from_skill_candidate_id",
            "to_skill_candidate_id",
            "relationship_type",
            "taxonomy_version",
            "review_status",
        ],
    ),
    "skill_source_mappings.jsonl": (
        "skill_source_mapping",
        "skill_source_mapping_id",
        [
            "skill_source_mapping_id",
            "skill_candidate_id",
            "curriculum_node_id",
            "mapping_type",
            "mapping_status",
            "mapping_version",
        ],
    ),
    "games.jsonl": (
        "game",
        "game_id",
        [
            "game_id",
            "game_version",
            "subject",
            "grade",
            "mechanic",
            "target_curriculum_node_ids",
            "required_modalities",
            "publication_status",
        ],
    ),
    "game_items.jsonl": (
        "game_item",
        "item_id",
        [
            "item_id",
            "item_version",
            "game_id",
            "subject",
            "language",
            "curriculum_node_ids",
            "response_format",
            "options",
            "correct_answers",
            "publication_status",
        ],
    ),
    "item_curriculum_mappings.jsonl": (
        "item_curriculum_mapping",
        "item_curriculum_mapping_id",
        [
            "item_curriculum_mapping_id",
            "item_id",
            "game_id",
            "curriculum_node_id",
            "relationship",
            "review_status",
        ],
    ),
    "item_skill_mappings.jsonl": (
        "item_skill_mapping",
        "item_skill_mapping_id",
        [
            "item_skill_mapping_id",
            "item_id",
            "skill_candidate_id",
            "relationship",
            "mapping_status",
            "mapping_version",
        ],
    ),
    "misconception_catalogue.jsonl": (
        "misconception",
        "misconception_code",
        [
            "misconception_code",
            "catalogue_version",
            "label_en",
            "approval_status",
        ],
    ),
    "structured_error_mappings.jsonl": (
        "structured_error_mapping",
        "structured_error_mapping_id",
        [
            "structured_error_mapping_id",
            "structured_error_code",
            "misconception_code",
            "mapping_status",
        ],
    ),
    "support_actions.jsonl": (
        "support_action",
        "action_id",
        [
            "action_id",
            "catalogue_version",
            "action_type",
            "content_asset_id",
            "languages",
            "skill_candidate_ids",
            "misconception_codes",
            "approved",
        ],
    ),
    "content_assets.jsonl": (
        "content_asset",
        "content_asset_id",
        [
            "content_asset_id",
            "content_version",
            "asset_type",
            "source_item_id",
            "languages",
            "publication_status",
        ],
    ),
    "verification_pair_candidates.jsonl": (
        "verification_pair_candidate",
        "verification_pair_id",
        [
            "verification_pair_id",
            "equivalence_version",
            "baseline_item_id",
            "verification_item_id",
            "skill_candidate_id",
            "approved",
            "review_status",
        ],
    ),
    "asset_requirements.jsonl": (
        "asset_requirement",
        "asset_requirement_id",
        [
            "asset_requirement_id",
            "game_id",
            "asset_type",
            "required_count",
            "language",
            "status",
        ],
    ),
    "qa_issues.jsonl": (
        "qa_issue",
        "qa_issue_id",
        [
            "qa_issue_id",
            "scope",
            "issue_type",
            "severity",
            "description",
            "status",
        ],
    ),
    "review_decisions.jsonl": (
        "review_decision",
        "review_decision_id",
        [
            "review_decision_id",
            "entity_type",
            "entity_id",
            "decision",
            "reviewer_id",
            "reviewed_at",
        ],
    ),
    "workbook_annotations.jsonl": (
        "workbook_annotation",
        "annotation_id",
        [
            "annotation_id",
            "source_workbook",
            "source_sheet",
            "source_row_number",
            "values",
            "instruction_authority",
        ],
    ),
}


def json_schema_for_values(values):
    type_names = set()
    list_values = []
    for value in values:
        if value is None:
            type_names.add("null")
        elif isinstance(value, bool):
            type_names.add("boolean")
        elif isinstance(value, int):
            type_names.add("integer")
        elif isinstance(value, float):
            type_names.add("number")
        elif isinstance(value, list):
            type_names.add("array")
            list_values.extend(value)
        elif isinstance(value, dict):
            type_names.add("object")
        else:
            type_names.add("string")
    schema = {"type": sorted(type_names)[0] if len(type_names) == 1 else sorted(type_names)}
    if "array" in type_names:
        schema["items"] = json_schema_for_values(list_values) if list_values else {}
    return schema


STATUS_RULES = {
    "draft": "Source-backed candidate that has not completed formal review.",
    "needs_review": "Requires named curriculum, language, answer, asset or equivalence review.",
    "published": "May be served only after every applicable release gate passes.",
}


def gamification_filenames(datasets):
    return [filename for filename in datasets if filename not in CURRICULUM_TRACE_FILES]


def schema_registry(datasets, filenames):
    definitions = {}
    files = {}
    for filename in filenames:
        entity, primary_key, required = FILE_ENTITY_SPECS[filename]
        rows = datasets[filename]
        field_names = sorted({field for row in rows for field in row} | set(required))
        properties = {}
        for field in field_names:
            observed = [row[field] for row in rows if field in row]
            properties[field] = json_schema_for_values(observed) if observed else {"type": "string"}
        properties[primary_key] = {"type": "string", "minLength": 1}
        definitions[entity] = {
            "type": "object",
            "required": required,
            "properties": properties,
            "additionalProperties": True,
        }
        files[filename] = {
            "entity": entity,
            "primary_key": primary_key,
            "record_schema": {"$ref": f"#/$defs/{entity}"},
        }
    return files, definitions


def schema_definition(datasets):
    files, definitions = schema_registry(datasets, gamification_filenames(datasets))
    return {
        "$schema": "https://json-schema.org/draft/2020-12/schema",
        "$id": "https://mindbridge.local/schemas/grade5-curriculum-game-dataset-v2.json",
        "title": "Grade 5 curriculum game dataset schema registry",
        "version": PACKAGE_VERSION,
        "curriculum_version": CURRICULUM_VERSION,
        "skill_taxonomy_version": SKILL_TAXONOMY_VERSION,
        "files": files,
        "$defs": definitions,
        "status_rules": STATUS_RULES,
    }


def curriculum_trace_schema(datasets):
    files, definitions = schema_registry(datasets, CURRICULUM_TRACE_FILES)
    return {
        "$schema": "https://json-schema.org/draft/2020-12/schema",
        "$id": "https://mindbridge.local/schemas/grade5-curriculum-trace-v1.json",
        "title": "Grade 5 curriculum trace schema registry",
        "version": CURRICULUM_TRACE_VERSION,
        "curriculum_version": CURRICULUM_VERSION,
        "files": files,
        "$defs": definitions,
        "status_rules": STATUS_RULES,
    }


def require_unique(rows, primary_key, entity_name):
    values = [row[primary_key] for row in rows]
    if len(values) != len(set(values)):
        raise ValueError(f"{entity_name} contains duplicate {primary_key} values")


def validate_datasets(datasets):
    checks = []
    for filename, (entity, primary_key, required) in FILE_ENTITY_SPECS.items():
        rows = datasets[filename]
        for index, row in enumerate(rows, start=1):
            missing = [field for field in required if field not in row]
            if missing:
                raise ValueError(f"{filename} row {index} is missing {missing}")
        require_unique(rows, primary_key, entity)
        checks.append({"check": f"{entity}_schema_and_unique_key", "status": "pass"})

    nodes = datasets["curriculum_nodes.jsonl"]
    games = datasets["games.jsonl"]
    items = datasets["game_items.jsonl"]
    sources = datasets["curriculum_sources.jsonl"]
    item_curriculum = datasets["item_curriculum_mappings.jsonl"]
    item_skills = datasets["item_skill_mappings.jsonl"]
    skills = datasets["skill_candidates.jsonl"]
    support_actions = datasets["support_actions.jsonl"]
    content_assets = datasets["content_assets.jsonl"]
    verification_pairs = datasets["verification_pair_candidates.jsonl"]

    node_ids = {row["curriculum_node_id"] for row in nodes}
    source_ids = {row["source_id"] for row in sources}
    game_ids = {row["game_id"] for row in games}
    item_ids = {row["item_id"] for row in items}
    skill_ids = {row["skill_candidate_id"] for row in skills}
    asset_ids = {row["content_asset_id"] for row in content_assets}

    if any(row["source_id"] not in source_ids for row in nodes):
        raise ValueError("curriculum node references an unknown source")
    if any(row["game_id"] not in game_ids for row in items):
        raise ValueError("game item references an unknown game")
    if any(
        curriculum_id not in node_ids
        for row in items
        for curriculum_id in row["curriculum_node_ids"]
    ):
        raise ValueError("game item references an unknown curriculum node")
    if any(
        row["item_id"] not in item_ids or row["curriculum_node_id"] not in node_ids
        for row in item_curriculum
    ):
        raise ValueError("item curriculum mapping contains an unknown reference")
    if any(
        row["item_id"] not in item_ids or row["skill_candidate_id"] not in skill_ids
        for row in item_skills
    ):
        raise ValueError("item skill mapping contains an unknown reference")
    if any(row["content_asset_id"] not in asset_ids for row in support_actions):
        raise ValueError("support action references an unknown content asset")
    if any(
        row["baseline_item_id"] not in item_ids
        or row["verification_item_id"] not in item_ids
        or row["skill_candidate_id"] not in skill_ids
        for row in verification_pairs
    ):
        raise ValueError("verification pair contains an unknown reference")

    for item in items:
        if item["response_format"] == "single_select" and len(
            item["options"]
        ) != len(set(map(str, item["options"]))):
            raise ValueError(f"{item['item_id']} contains duplicate options")
        if item["response_format"] == "single_select" and any(
            answer not in item["options"] for answer in item["correct_answers"]
        ):
            raise ValueError(f"{item['item_id']} has a correct answer outside its options")
    checks.extend(
        [
            {"check": "all_references_resolve", "status": "pass"},
            {"check": "single_select_answers_are_valid", "status": "pass"},
            {"check": "option_lists_have_no_duplicates", "status": "pass"},
        ]
    )
    return {
        "schema_version": PACKAGE_VERSION,
        "status": "pass",
        "checks": checks,
    }


def build_readiness_report(datasets, validation_report):
    nodes = datasets["curriculum_nodes.jsonl"]
    items = datasets["game_items.jsonl"]
    item_skill_mappings = datasets["item_skill_mappings.jsonl"]
    support_actions = datasets["support_actions.jsonl"]
    pairs = datasets["verification_pair_candidates.jsonl"]
    qa_issues = datasets["qa_issues.jsonl"]
    mapped_item_ids = {row["item_id"] for row in item_skill_mappings}
    return {
        "package_version": PACKAGE_VERSION,
        "overall_status": "engineering_ready_research_and_learner_delivery_blocked",
        "summary": {
            "curriculum_sources": len(datasets["curriculum_sources.jsonl"]),
            "curriculum_nodes": len(nodes),
            "curriculum_terms": len(datasets["curriculum_terms.jsonl"]),
            "games": len(datasets["games.jsonl"]),
            "game_items": len(items),
            "item_curriculum_mappings": len(
                datasets["item_curriculum_mappings.jsonl"]
            ),
            "skill_candidates": len(datasets["skill_candidates.jsonl"]),
            "items_with_skill_candidate_mapping": len(mapped_item_ids),
            "items_without_skill_candidate_mapping": len(items) - len(mapped_item_ids),
            "answer_verified_items": sum(bool(row["answer_verified"]) for row in items),
            "language_reviewed_items": sum(bool(row["language_reviewed"]) for row in items),
            "published_items": sum(row["publication_status"] == "published" for row in items),
            "approved_curriculum_nodes": sum(
                row["verification_status"] == "approved" for row in nodes
            ),
            "support_action_candidates": len(support_actions),
            "approved_support_actions": sum(bool(row["approved"]) for row in support_actions),
            "verification_pair_candidates": len(pairs),
            "approved_verification_pairs": sum(bool(row["approved"]) for row in pairs),
            "items_without_verification_pair_candidate": len(items)
            - len({row["baseline_item_id"] for row in pairs}),
            "structured_error_mappings": len(
                datasets["structured_error_mappings.jsonl"]
            ),
            "formal_review_decisions": len(datasets["review_decisions.jsonl"]),
            "audio_required_items": sum(
                "audio" in row["required_asset_types"] for row in items
            ),
            "uncreated_required_audio_assets": sum(
                row["required_count"]
                for row in datasets["asset_requirements.jsonl"]
                if row["asset_type"] == "audio" and row["status"] != "created"
            ),
            "open_high_severity_qa_issues": sum(
                row["severity"] == "high" and row["status"] == "open"
                for row in qa_issues
            ),
            "open_medium_severity_qa_issues": sum(
                row["severity"] == "medium" and row["status"] == "open"
                for row in qa_issues
            ),
        },
        "gates": {
            "package_integrity": {
                "status": "ready" if validation_report["status"] == "pass" else "blocked",
                "reason": "All generated schemas, identifiers and cross-references validate.",
            },
            "game_prototyping": {
                "status": "ready_with_restrictions",
                "reason": "Structurally playable content exists, but publication and asset reviews remain open.",
            },
            "curriculum_and_skill_approval": {
                "status": "blocked",
                "reason": "Gamification skill candidates have no formal curriculum approvals.",
            },
            "misconception_inference": {
                "status": "blocked",
                "reason": "Only the unknown/other fallback exists; no expert-approved structured error mappings exist.",
            },
            "support_selection": {
                "status": "blocked",
                "reason": "Support actions and assets are review candidates and none are approved for eligibility.",
            },
            "independence_gain": {
                "status": "blocked",
                "reason": "Verification pairs are candidates only; unseen-item and answer-overlap evidence are runtime or review requirements.",
            },
            "child_facing_delivery": {
                "status": "blocked",
                "reason": "No content is published and curriculum, language, asset and ethics gates remain open.",
            },
        },
    }


def package_readme():
    return f"""# Grade 5 curriculum game dataset v{PACKAGE_VERSION}

This is the canonical package. It contains loss-aware source evidence, explicit curriculum and
skill-candidate mappings, model-facing support and verification candidates, QA data, formal
schemas, integrity validation and a machine-readable readiness report.

All skill, misconception, support and verification records remain candidates unless their explicit
approval fields say otherwise. Source workbook notes and instructions are stored as reference text
with `instruction_authority: none`; they are never executed as package-building instructions.

The package is suitable for engineering and review workflows. It is not approved for child-facing
delivery, research claims or learned-model activation. See `readiness_report.json` for current gates.

Curriculum sources, nodes and terms are not stored here. They are shared with the Adaptive Tutor
in `data/curriculum/{CURRICULUM_TRACE_PACKAGE}/v1/`. `curriculum_node_id` and `source_id` values
in this package refer to that shared package, whose version and manifest hash are recorded in
`manifest.json` under `curriculum_trace_package`.
"""


def curriculum_trace_readme():
    return f"""# Grade 5 curriculum trace v{CURRICULUM_TRACE_VERSION}

Shared curriculum facts used by both the Adaptive Tutor and the Gamification component, and the
seed for the shared curriculum RAG (`docs/curriculum/shared-rag-joint-plan.md`).

- `curriculum_sources.jsonl`: Grade 5 textbooks, teacher guides and activity books that were traced.
- `curriculum_nodes.jsonl`: competencies, competency levels, lessons, learning outcomes, content
  points and source activities, with Sinhala text, page ranges and parent links.
- `curriculum_terms.jsonl`: subject terms with Sinhala and English labels and source pages.

Every record is a draft trace transcribed by visual reading or OCR. Nothing here is reviewed or
approved, and no record may reach a learner until a named reviewer approves it. Source workbook
notes are reference text only and are never executed as instructions.

This package is generated by `tools/gamification/datasets/build_grade5_curriculum_game_dataset.py`
together with the gamification dataset. Keep `curriculum_node_id` and `source_id` values stable;
add a new version folder for incompatible changes. The shared source registry in
`data/curriculum/registry/v1/` is generated from `curriculum_sources.jsonl`.
"""


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Build the versioned Grade 5 curriculum-to-game dataset."
    )
    parser.add_argument(
        "--source-dir",
        type=Path,
        default=Path(os.environ[SOURCE_DIR_ENV]) if SOURCE_DIR_ENV in os.environ else None,
        help=(
            "Directory containing the four syllabus trace workbooks and the game dataset "
            f"master workbook (or set {SOURCE_DIR_ENV})."
        ),
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=DEFAULT_OUTPUT_DIR,
        help=f"Gamification package destination (default: {DEFAULT_OUTPUT_DIR}).",
    )
    parser.add_argument(
        "--curriculum-output-dir",
        type=Path,
        default=DEFAULT_CURRICULUM_OUTPUT_DIR,
        help=f"Shared curriculum trace destination (default: {DEFAULT_CURRICULUM_OUTPUT_DIR}).",
    )
    parser.add_argument(
        "--repackage-existing",
        action="store_true",
        help=(
            "Rewrite both packages from their existing records instead of the source workbooks. "
            "Use this to regenerate schemas, reports and manifests without the workbooks."
        ),
    )
    args = parser.parse_args()
    if args.source_dir is None and not args.repackage_existing:
        parser.error(
            f"--source-dir is required unless {SOURCE_DIR_ENV} is set or --repackage-existing is used"
        )
    return args


def validate_source_dir(source_dir: Path) -> None:
    required = [
        source_dir / f"grade5_{subject}_syllabus_trace_v1.xlsx"
        for subject in SUBJECTS
    ]
    required.append(source_dir / MASTER_WORKBOOK_NAME)
    missing = [path for path in required if not path.is_file()]
    if missing:
        formatted = "\n".join(f"- {path}" for path in missing)
        raise FileNotFoundError(f"Missing required Grade 5 workbooks:\n{formatted}")


def describe_source_workbooks(source_dir):
    source_workbooks = [
        *(f"grade5_{subject}_syllabus_trace_v1.xlsx" for subject in SUBJECTS),
        MASTER_WORKBOOK_NAME,
    ]
    descriptions = {}
    for filename in source_workbooks:
        path = source_dir / filename
        descriptions[filename] = {
            "prepared_by": "hareeshagunasekara",
            "bytes": path.stat().st_size,
            "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
        }
    return descriptions


def read_jsonl(path):
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines()]


def load_existing_packages(output_dir, curriculum_output_dir):
    datasets = {}
    for filename in FILE_ENTITY_SPECS:
        directory = curriculum_output_dir if filename in CURRICULUM_TRACE_FILES else output_dir
        datasets[filename] = read_jsonl(directory / filename)
    manifest = json.loads((output_dir / "manifest.json").read_text(encoding="utf-8"))
    return datasets, manifest["source_workbooks"]


def file_descriptions(directory, filenames):
    descriptions = {}
    for filename in filenames:
        payload = (directory / filename).read_bytes()
        descriptions[filename] = {
            "bytes": len(payload),
            "sha256": hashlib.sha256(payload).hexdigest(),
        }
    return descriptions


def write_json(path, payload):
    path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
        newline="\n",
    )


def write_curriculum_trace_package(datasets, source_workbooks, curriculum_output_dir):
    curriculum_output_dir.mkdir(parents=True, exist_ok=True)
    for filename in CURRICULUM_TRACE_FILES:
        jsonl_write(curriculum_output_dir / filename, datasets[filename])
    write_json(curriculum_output_dir / "schema.json", curriculum_trace_schema(datasets))
    (curriculum_output_dir / "README.md").write_text(
        curriculum_trace_readme(), encoding="utf-8", newline="\n"
    )
    manifest = {
        "package": CURRICULUM_TRACE_PACKAGE,
        "version": CURRICULUM_TRACE_VERSION,
        "curriculum_version": CURRICULUM_VERSION,
        "ownership": "shared curriculum (Adaptive Tutor and Gamification)",
        "review_status": "draft trace; not reviewed or approved",
        "counts": {
            FILE_ENTITY_SPECS[filename][0]: len(datasets[filename])
            for filename in CURRICULUM_TRACE_FILES
        },
        "source_workbooks": {
            filename: description
            for filename, description in source_workbooks.items()
            if filename != MASTER_WORKBOOK_NAME
        },
        "files": file_descriptions(
            curriculum_output_dir, (*CURRICULUM_TRACE_FILES, "schema.json", "README.md")
        ),
    }
    manifest_path = curriculum_output_dir / "manifest.json"
    manifest_path.write_text(
        json.dumps(manifest, indent=2) + "\n", encoding="utf-8", newline="\n"
    )
    return hashlib.sha256(manifest_path.read_bytes()).hexdigest()


def write_packages(
    datasets,
    validation_report,
    readiness_report,
    source_workbooks,
    output_dir,
    curriculum_output_dir,
):
    curriculum_manifest_sha256 = write_curriculum_trace_package(
        datasets, source_workbooks, curriculum_output_dir
    )
    output_dir.mkdir(parents=True, exist_ok=True)
    package_files = gamification_filenames(datasets)
    for filename in package_files:
        jsonl_write(output_dir / filename, datasets[filename])
    write_json(output_dir / "schema.json", schema_definition(datasets))
    write_json(output_dir / "validation_report.json", validation_report)
    write_json(output_dir / "readiness_report.json", readiness_report)
    (output_dir / "README.md").write_text(
        package_readme(), encoding="utf-8", newline="\n"
    )

    manifest = {
        "package": "grade5-curriculum-game-dataset",
        "component_owner": "hareeshagunasekara",
        "version": PACKAGE_VERSION,
        "curriculum_version": CURRICULUM_VERSION,
        "skill_taxonomy_version": SKILL_TAXONOMY_VERSION,
        "misconception_catalogue_version": MISCONCEPTION_CATALOGUE_VERSION,
        "support_catalogue_version": SUPPORT_CATALOGUE_VERSION,
        "equivalence_version": EQUIVALENCE_VERSION,
        "curriculum_trace_package": {
            "package": CURRICULUM_TRACE_PACKAGE,
            "version": CURRICULUM_TRACE_VERSION,
            "path": Path(os.path.relpath(curriculum_output_dir, output_dir)).as_posix(),
            "manifest_sha256": curriculum_manifest_sha256,
        },
        "counts": {
            FILE_ENTITY_SPECS[filename][0]: len(datasets[filename]) for filename in package_files
        },
        "source_workbooks": source_workbooks,
        "files": file_descriptions(
            output_dir,
            (
                *package_files,
                "schema.json",
                "validation_report.json",
                "readiness_report.json",
                "README.md",
            ),
        ),
    }
    (output_dir / "manifest.json").write_text(
        json.dumps(manifest, indent=2) + "\n", encoding="utf-8", newline="\n"
    )
    return manifest


def build_datasets_from_workbooks(source_dir):
    validate_source_dir(source_dir)
    RNG.seed(5323)
    nodes, raw = collect_curriculum(source_dir)
    raw_games = build_game_catalog(nodes)
    raw_items = build_items(raw_games, raw)
    games, items = normalize_games_and_items(raw_games, raw_items)
    (
        item_curriculum_mappings,
        master_qa_issues,
        asset_requirements,
        workbook_annotations,
    ) = collect_master_workbook(source_dir)
    curriculum_sources, curriculum_terms, source_issues = collect_source_entities(
        raw, source_dir
    )
    source_issue_by_id = {issue["qa_issue_id"]: issue for issue in source_issues}
    for issue in master_qa_issues:
        source_issue = source_issue_by_id.get(issue["qa_issue_id"])
        if source_issue:
            issue["source_trace_record"] = source_issue["source_record"]
            issue["source_trace_workbook"] = source_issue["source_workbook"]
            issue["source_trace_sheet"] = source_issue["source_sheet"]
            issue["source_trace_row_number"] = source_issue["source_row_number"]

    skills, skill_relationships, skill_source_mappings = build_skill_entities(nodes)
    item_skill_mappings = build_item_skill_mappings(
        items, item_curriculum_mappings, nodes, skills
    )
    support_actions, content_assets = build_support_entities(items, item_skill_mappings)
    verification_pairs = build_verification_pair_candidates(items)
    qa_issues = master_qa_issues + generated_qa_issues(
        items, item_skill_mappings, support_actions, verification_pairs
    )

    misconception_catalogue = [
        {
            "misconception_code": "unknown/other",
            "catalogue_version": MISCONCEPTION_CATALOGUE_VERSION,
            "label_primary": "නොදන්නා හෝ වෙනත්",
            "label_en": "Unknown or other",
            "description": "Conservative fallback when no expert-approved diagnosis is available.",
            "applicable_skill_candidate_ids": [],
            "approval_status": "engineering_fallback_only",
            "child_facing_diagnosis_approved": False,
        }
    ]
    datasets = {
        "curriculum_sources.jsonl": curriculum_sources,
        "curriculum_nodes.jsonl": nodes,
        "curriculum_terms.jsonl": curriculum_terms,
        "skill_candidates.jsonl": skills,
        "skill_relationships.jsonl": skill_relationships,
        "skill_source_mappings.jsonl": skill_source_mappings,
        "games.jsonl": games,
        "game_items.jsonl": items,
        "item_curriculum_mappings.jsonl": item_curriculum_mappings,
        "item_skill_mappings.jsonl": item_skill_mappings,
        "misconception_catalogue.jsonl": misconception_catalogue,
        "structured_error_mappings.jsonl": [],
        "support_actions.jsonl": support_actions,
        "content_assets.jsonl": content_assets,
        "verification_pair_candidates.jsonl": verification_pairs,
        "asset_requirements.jsonl": asset_requirements,
        "qa_issues.jsonl": qa_issues,
        "review_decisions.jsonl": [],
        "workbook_annotations.jsonl": workbook_annotations,
    }
    return datasets, describe_source_workbooks(source_dir)


def main() -> None:
    args = parse_args()
    output_dir = args.output_dir.expanduser().resolve()
    curriculum_output_dir = args.curriculum_output_dir.expanduser().resolve()
    if args.repackage_existing:
        datasets, source_workbooks = load_existing_packages(output_dir, curriculum_output_dir)
    else:
        datasets, source_workbooks = build_datasets_from_workbooks(
            args.source_dir.expanduser().resolve()
        )
    validation_report = validate_datasets(datasets)
    readiness_report = build_readiness_report(datasets, validation_report)
    manifest = write_packages(
        datasets,
        validation_report,
        readiness_report,
        source_workbooks,
        output_dir,
        curriculum_output_dir,
    )
    print(
        json.dumps(
            {
                "counts": manifest["counts"],
                "curriculum_trace_package": manifest["curriculum_trace_package"],
                "readiness": readiness_report["overall_status"],
            }
        )
    )


if __name__ == "__main__":
    main()
